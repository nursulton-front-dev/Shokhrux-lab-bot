"""Adversarial admin regression tests; no external Telegram or production DB."""
import asyncio
import datetime
import threading
import time
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pytest
from aiogram.dispatcher.event.bases import UNHANDLED
from openpyxl import load_workbook

from bot.database.models import Subscription, User
from bot.filters.admin import IsAdmin
from bot.handlers import admin


@pytest.fixture
def callback():
    return SimpleNamespace(
        data="", from_user=SimpleNamespace(id=900),
        answer=AsyncMock(),
        message=SimpleNamespace(html_text="Payment", answer=AsyncMock(),
                                edit_text=AsyncMock(), edit_reply_markup=AsyncMock()),
    )


@pytest.mark.asyncio
@pytest.mark.parametrize("data", [
    "admin_conf_1_42", "admin_rej_1_42", "adm_user_link_42", "adm_user_kick_42",
    "adm_user_ext_42", "adm_ticket_reply_1", "users_excel_download", "vip_excel_download",
])
async def test_every_admin_callback_rejects_non_admin(monkeypatch, callback, data):
    monkeypatch.setattr(admin.config, "admin_ids", "900")
    monkeypatch.setattr(admin.config, "admin_id", None)
    monkeypatch.setattr(admin.config, "support_id", None)
    callback.from_user.id = 12345
    callback.data = data
    session, bot = AsyncMock(), AsyncMock()
    result = await admin.router.propagate_event(
        update_type="callback_query", event=callback, session=session, bot=bot,
    )
    assert result is UNHANDLED
    session.scalar.assert_not_awaited()
    bot.send_message.assert_not_awaited()
    bot.create_chat_invite_link.assert_not_awaited()


@pytest.mark.asyncio
async def test_admin_filter_handles_missing_message_sender(monkeypatch):
    monkeypatch.setattr(admin.config, "admin_ids", "900")
    assert await IsAdmin()(SimpleNamespace(from_user=None)) is False
    assert await IsAdmin()(SimpleNamespace(from_user=SimpleNamespace(id=900))) is True


@pytest.mark.parametrize("raw", ["", "-50", "1e9", "nan", "9999999999999999999999999", "9223372036854775808", "１２３", "0"])
def test_positive_id_rejects_invalid_database_ids(raw):
    assert admin._positive_id(raw) is None


@pytest.mark.asyncio
@pytest.mark.parametrize("handler,prefix", [
    (admin.adm_user_extend, "adm_user_ext_"),
    (admin.adm_user_kick, "adm_user_kick_"),
    (admin.adm_user_link, "adm_user_link_"),
])
async def test_malformed_admin_callback_does_no_db_work(callback, handler, prefix):
    callback.data = prefix + "1e9"
    session, bot = AsyncMock(), AsyncMock()
    kwargs = {"callback": callback, "session": session}
    if handler is not admin.adm_user_extend:
        kwargs["bot"] = bot
    await handler(**kwargs)
    session.scalar.assert_not_awaited()
    callback.answer.assert_awaited_once()


@pytest.mark.asyncio
async def test_reject_uses_durable_payment_owner_not_callback_owner(monkeypatch, callback):
    callback.data = "admin_rej_12_999"
    reject = AsyncMock(return_value=42)
    monkeypatch.setattr(admin, "reject_pending_payment", reject)
    session, bot = AsyncMock(), AsyncMock()
    session.scalar.return_value = SimpleNamespace(language="ru")
    await admin.admin_reject_payment(callback, session, bot)
    reject.assert_awaited_once_with(12, session)
    assert bot.send_message.await_args.kwargs["chat_id"] == 42
    session.commit.assert_not_awaited()  # service owns the status transition


@pytest.mark.asyncio
async def test_reject_already_completed_never_notifies(monkeypatch, callback):
    callback.data = "admin_rej_12_42"
    monkeypatch.setattr(admin, "reject_pending_payment", AsyncMock(return_value=None))
    bot = AsyncMock()
    await admin.admin_reject_payment(callback, AsyncMock(), bot)
    bot.send_message.assert_not_awaited()


def export_rows(name='=HYPERLINK("https://example.invalid","open")\x00'):
    now = datetime.datetime.now(datetime.timezone.utc)
    user = admin.ExportUser(42, "some_user", name, "+998901234567", "ru", 0, None, now)
    sub = admin.ExportSubscription(42, "active", 6, now + datetime.timedelta(days=180))
    pay = admin.ExportPayment(1, 42, "completed", 6, 300000, 0, now)
    return user, sub, pay, now.replace(tzinfo=None)


def test_all_excel_sheets_serialize_untrusted_data_as_text():
    user, sub, pay, now = export_rows()
    buffer = admin._render_styled_excel([user], [sub], [pay], now)
    workbook = load_workbook(buffer, data_only=False)
    for name in ["База пользователей", "VIP Клиенты (6 мес)", "Транзакции"]:
        cell = workbook[name]["C2"]
        assert cell.data_type == "s"
        assert cell.value.startswith("'=HYPERLINK")
        assert "\x00" not in cell.value
    assert workbook["База пользователей"]["D2"].value == "'+998901234567"
    assert all(cell.data_type != "f" for ws in workbook for row in ws for cell in row)
    workbook.close()
    vip = load_workbook(admin._render_vip_excel([(sub, user)], now), data_only=False)
    assert vip.active["C2"].data_type == "s"
    assert vip.active["C2"].value.startswith("'=HYPERLINK")
    vip.close()


@pytest.mark.asyncio
@pytest.mark.parametrize("export_name,render_name", [
    ("generate_styled_excel", "_render_styled_excel"),
    ("generate_vip_excel", "_render_vip_excel"),
])
async def test_excel_render_and_save_do_not_block_event_loop(monkeypatch, export_name, render_name):
    monkeypatch.setattr(admin, "_excel_lock", asyncio.Lock())
    monkeypatch.setattr(admin, "_get_active_vip_clients", AsyncMock(return_value=[]))
    session = AsyncMock()
    session.scalars.return_value = SimpleNamespace(all=lambda: [])
    main_thread = threading.get_ident()
    worker_threads = []
    original = getattr(admin, render_name)

    def slow_render(*args):
        worker_threads.append(threading.get_ident())
        time.sleep(0.08)  # stands in for CPU/XML/zip writing
        return original(*args)

    monkeypatch.setattr(admin, render_name, slow_render)
    task = asyncio.create_task(getattr(admin, export_name)(session))
    heartbeats = 0
    while not task.done():
        await asyncio.sleep(0.005)
        heartbeats += 1
    workbook = load_workbook(await task)
    workbook.close()
    assert worker_threads and all(tid != main_thread for tid in worker_threads)
    assert heartbeats >= 5


@pytest.mark.asyncio
async def test_vip_html_name_cannot_inject_or_break_markup(monkeypatch):
    user, sub, _, _ = export_rows("<b>Injected & name</b>")
    monkeypatch.setattr(admin, "_get_active_vip_clients", AsyncMock(return_value=[(sub, user)]))
    event = SimpleNamespace(answer=AsyncMock(), message=SimpleNamespace(answer=AsyncMock()))
    await admin.msg_admin_vip_clients(event, AsyncMock())
    text = event.message.answer.await_args.args[0]
    assert "&lt;b&gt;Injected &amp; name&lt;/b&gt;" in text
    assert "<b>Injected" not in text


@pytest.mark.asyncio
async def test_kick_failure_retains_durable_retry_marker(monkeypatch, callback):
    callback.data = "adm_user_kick_42"
    sub = Subscription(id=1, user_id=42, status="active", tariff_months=6,
                       expires_at=datetime.datetime.now(datetime.timezone.utc),
                       invite_link="https://t.me/+old", vip_invite_link=None)
    session = AsyncMock()
    session.scalar.return_value = User(telegram_id=42)
    session.scalars.return_value = SimpleNamespace(all=lambda: [sub])
    monkeypatch.setattr(admin.config, "channel_id", -100123)
    monkeypatch.setattr(admin, "revoke_invite", AsyncMock(side_effect=TimeoutError("Telegram timeout")))
    monkeypatch.setattr(admin, "kick_member", AsyncMock())
    await admin.adm_user_kick(callback, session, AsyncMock())
    assert sub.status == "revocation_pending"
    session.commit.assert_awaited_once()
    session.rollback.assert_awaited_once()
    admin.kick_member.assert_not_awaited()


@pytest.mark.asyncio
async def test_manual_invite_requires_owner_checked_join_request(monkeypatch, callback):
    callback.data = "adm_user_link_42"
    expiry = datetime.datetime.now(datetime.timezone.utc) + datetime.timedelta(days=30)
    sub = Subscription(id=1, user_id=42, status="active", tariff_months=1,
                       expires_at=expiry, invite_link=None)
    session = AsyncMock()
    session.scalar.side_effect = [User(telegram_id=42), sub]
    bot = AsyncMock()
    bot.create_chat_invite_link.return_value = SimpleNamespace(invite_link="https://t.me/+new")
    monkeypatch.setattr(admin.config, "channel_id", -100123)
    await admin.adm_user_link(callback, session, bot)
    kwargs = bot.create_chat_invite_link.await_args.kwargs
    assert kwargs["creates_join_request"] is True
    assert kwargs["expire_date"] == expiry
    assert "member_limit" not in kwargs
    assert sub.invite_link == "https://t.me/+new"
    session.commit.assert_awaited_once()
