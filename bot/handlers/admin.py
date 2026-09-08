import datetime
import os
import asyncio
import logging
from io import BytesIO
from dataclasses import dataclass
from html import escape
from typing import Callable, TypeVar

from aiogram import Router, F, Bot
from aiogram.types import (
    Message, 
    CallbackQuery, 
    InlineKeyboardMarkup, 
    InlineKeyboardButton, 
    BufferedInputFile,
    ReplyKeyboardMarkup, 
    KeyboardButton, 
    ReplyKeyboardRemove
)
from aiogram.filters import Command, CommandStart
from aiogram.fsm.context import FSMContext
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, exists, or_, desc

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from bot.database.models import User, Subscription, Payment, Ticket, CashbackTransaction
from bot.filters.admin import IsAdmin
from bot.states.admin import AdminBroadcast, AdminUserSearch
from bot.states.support import AdminTicketState
from bot.config import config
from bot.services.rahmat import process_successful_payment, reject_pending_payment
from bot.services.background_tasks import create_background_task
from bot.services.telegram_rate_limit import kick_member, revoke_invite
from bot import texts

logger = logging.getLogger(__name__)

# Snapshot loaded scalar columns before entering a worker: AsyncSession and ORM
# state never cross the thread boundary. One render at a time bounds memory use.
_excel_lock = asyncio.Lock()
_Result = TypeVar("_Result")


@dataclass(frozen=True, slots=True)
class ExportUser:
    telegram_id: int
    username: str | None
    full_name: str | None
    phone_number: str | None
    language: str | None
    balance: int
    referred_by: int | None
    created_at: datetime.datetime | None


@dataclass(frozen=True, slots=True)
class ExportSubscription:
    user_id: int
    status: str
    tariff_months: int
    expires_at: datetime.datetime | None


@dataclass(frozen=True, slots=True)
class ExportPayment:
    id: int
    user_id: int
    status: str
    tariff_months: int
    amount: int
    cashback_applied: int
    created_at: datetime.datetime | None


def _export_user(user: User) -> ExportUser:
    return ExportUser(user.telegram_id, user.username, user.full_name,
                      user.phone_number, user.language, user.balance or 0,
                      user.referred_by, user.created_at)


def _export_subscription(sub: Subscription) -> ExportSubscription:
    return ExportSubscription(sub.user_id, sub.status, sub.tariff_months, sub.expires_at)


def _export_payment(payment: Payment) -> ExportPayment:
    return ExportPayment(payment.id, payment.user_id, payment.status,
                         payment.tariff_months, payment.amount,
                         payment.cashback_applied or 0, payment.created_at)


def _excel_text(value: str | None) -> str:
    """Store untrusted content as inert text, including illegal XML characters."""
    value = ILLEGAL_CHARACTERS_RE.sub("", value or "-")[:32766]
    if value.lstrip().startswith(("=", "+", "-", "@")):
        value = "'" + value
    return value


async def _render_in_thread(render: Callable[..., _Result], *args: object) -> _Result:
    task = asyncio.create_task(asyncio.to_thread(render, *args))
    try:
        return await asyncio.shield(task)
    except asyncio.CancelledError:
        # to_thread cannot stop a running worker. Keep the export lock held until
        # it really exits, even when the requesting update is cancelled.
        try:
            await task
        finally:
            raise


def _positive_id(raw: str) -> int | None:
    if not raw.isascii() or not raw.isdecimal() or len(raw) > 19:
        return None
    value = int(raw)
    return value if 0 < value <= 2**63 - 1 else None


router = Router()
router.message.filter(IsAdmin())
router.callback_query.filter(IsAdmin())

def get_channel_id() -> int:
    return config.channel_id or int(os.getenv("CHANNEL_ID", "-1001234567890"))

def to_naive_utc(dt: datetime.datetime | None) -> datetime.datetime | None:
    if dt is None:
        return None
    if getattr(dt, 'tzinfo', None) is not None:
        return dt.astimezone(datetime.timezone.utc).replace(tzinfo=None)
    return dt

# ======================== 🎛️ ДВУХУРОВНЕВОЕ АДМИН-МЕНЮ ========================

async def get_admin_main_kb_async(session: AsyncSession):
    stmt = select(func.count()).select_from(Ticket).where(Ticket.status == "open")
    open_tickets = (await session.scalar(stmt)) or 0
    return ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text=f"📥 Заявки / Тикеты ({open_tickets} шт.)")],
            [KeyboardButton(text="👥 Пользователи (База Excel)"), KeyboardButton(text="👑 VIP Клиенты (6 мес)")],
            [KeyboardButton(text="🔎 Поиск клиента"), KeyboardButton(text="📈 Отчёты и Аналитика")],
            [KeyboardButton(text="📢 Рассылка сообщений")]
        ],
        resize_keyboard=True
    )

def get_admin_analytics_kb():
    return ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text="📅 Касса за сегодня"), KeyboardButton(text="📉 Отчёт за текущий месяц")],
            [KeyboardButton(text="🌐 За всё время (LTV)"), KeyboardButton(text="🔄 Продления (Retention)")],
            [KeyboardButton(text="📥 Скачать ПОЛНЫЙ Excel"), KeyboardButton(text="◀️ Назад в админку")]
        ],
        resize_keyboard=True
    )

@router.message(Command("admin"))
@router.message(CommandStart())
async def cmd_admin(message: Message, state: FSMContext, session: AsyncSession):
    await state.clear()
    kb = await get_admin_main_kb_async(session)
    await message.answer("👑 <b>Центр управления ботом</b>\nВыберите раздел в меню ниже 👇", reply_markup=kb)

@router.message(F.text.in_(["👑 Главное меню", "◀️ Назад в админку"]))
async def msg_admin_main(message: Message, state: FSMContext, session: AsyncSession):
    await state.clear()
    kb = await get_admin_main_kb_async(session)
    await message.answer("👑 <b>Главное меню админ-панели</b>\nВыберите раздел:", reply_markup=kb)

@router.callback_query(F.data == "admin_main")
async def cb_admin_main(callback: CallbackQuery, state: FSMContext, session: AsyncSession):
    await state.clear()
    try:
        await callback.message.delete()
    except Exception:
        pass
    kb = await get_admin_main_kb_async(session)
    await callback.message.answer("👑 <b>Главное меню админ-панели</b>\nВыберите раздел:", reply_markup=kb)
    await callback.answer()

@router.message(F.text.in_(["📈 Отчёты и Аналитика", "📊 Аналитика и Отчёты", "📈 Отчеты и Аналитика"]))
async def msg_admin_analytics_menu(message: Message, state: FSMContext):
    await state.clear()
    kb = get_admin_analytics_kb()
    await message.answer("📈 <b>Раздел аналитики и отчётов</b>\nВыберите интересующий срез данных или скачайте ПОЛНЫЙ Excel 👇", reply_markup=kb)

@router.message(F.text == "📅 Касса за сегодня")
async def msg_an_today(message: Message, session: AsyncSession):
    now = datetime.datetime.utcnow()
    start_today = datetime.datetime(now.year, now.month, now.day)
    
    payments = (await session.execute(select(Payment).where(Payment.status == "completed"))).scalars().all()
    today_payments = [p for p in payments if to_naive_utc(p.created_at) and to_naive_utc(p.created_at) >= start_today]
    
    revenue = sum(p.amount for p in today_payments)
    cnt = len(today_payments)
    
    text = (
        f"📅 <b>КАССА ЗА СЕГОДНЯ</b> ({now.strftime('%d.%m.%Y')})\n\n"
        f"💰 Выручка: <b>{revenue:,} UZS</b>\n"
        f"💳 Оплаченных заказов: <b>{cnt} шт.</b>"
    )
    await message.answer(text)

@router.message(F.text.in_(["📉 Отчёт за текущий месяц", "📈 Отчёт за текущий месяц"]))
async def msg_an_month(message: Message, session: AsyncSession):
    now = datetime.datetime.utcnow()
    start_month = datetime.datetime(now.year, now.month, 1)
    
    payments = (await session.execute(select(Payment).where(Payment.status == "completed"))).scalars().all()
    month_payments = [p for p in payments if to_naive_utc(p.created_at) and to_naive_utc(p.created_at) >= start_month]
    
    revenue = sum(p.amount for p in month_payments)
    cnt = len(month_payments)
    
    text = (
        f"📉 <b>ОТЧЁТ ЗА ТЕКУЩИЙ МЕСЯЦ</b> ({now.strftime('%m.%Y')})\n\n"
        f"💰 Выручка за месяц: <b>{revenue:,} UZS</b>\n"
        f"💳 Оплаченных заказов: <b>{cnt} шт.</b>"
    )
    await message.answer(text)

@router.message(F.text == "🌐 За всё время (LTV)")
async def msg_an_alltime(message: Message, session: AsyncSession):
    stmt = select(func.sum(Payment.amount)).where(Payment.status == "completed")
    revenue = (await session.scalar(stmt)) or 0
    
    stmt_cnt = select(func.count(Payment.id)).where(Payment.status == "completed")
    cnt = (await session.scalar(stmt_cnt)) or 0
    
    u_count = (await session.scalar(select(func.count()).select_from(User))) or 0
    
    text = (
        f"🌐 <b>АНАЛИТИКА ЗА ВСЁ ВРЕМЯ (LTV)</b>\n\n"
        f"👥 Всего зарегистрировано: <b>{u_count} чел.</b>\n"
        f"💰 Общая выручка: <b>{revenue:,} UZS</b>\n"
        f"💳 Всего успешных транзакций: <b>{cnt} шт.</b>"
    )
    await message.answer(text)

@router.message(F.text == "🔄 Продления (Retention)")
async def msg_an_retention(message: Message, session: AsyncSession):
    stmt_buys = select(Payment.user_id, func.count(Payment.id).label("buy_count")).where(Payment.status == "completed").group_by(Payment.user_id)
    res = await session.execute(stmt_buys)
    rows = res.all()
    
    total_buyers = len(rows)
    repeat_buyers = sum(1 for r in rows if r.buy_count > 1)
    
    retention_rate = (repeat_buyers / total_buyers * 100) if total_buyers > 0 else 0.0
    
    text = (
        f"🔄 <b>АНАЛИЗ ПРОДЛЕНИЙ (RETENTION)</b>\n\n"
        f"👥 Покупателей всего: <b>{total_buyers}</b>\n"
        f"🔁 Повторных покупок: <b>{repeat_buyers}</b>\n"
        f"📈 Коэффициент удержания (Retention): <b>{retention_rate:.1f}%</b>"
    )
    await message.answer(text)

@router.callback_query(F.data == "noop")
async def cb_noop(callback: CallbackQuery):
    await callback.answer()

# ======================== 👑 VIP КЛИЕНТЫ (6 МЕСЯЦЕВ) ========================

async def _get_active_vip_clients(
    session: AsyncSession,
) -> list[tuple[ExportSubscription, ExportUser]]:
    now = datetime.datetime.now(datetime.timezone.utc)
    rows = (await session.execute(
        select(Subscription, User)
        .join(User, User.telegram_id == Subscription.user_id)
        .where(Subscription.status == "active", Subscription.tariff_months == 6,
               Subscription.expires_at > now)
        .order_by(Subscription.expires_at)
    )).all()
    return [(_export_subscription(sub), _export_user(user)) for sub, user in rows]


@router.message(F.text.in_(["👑 VIP Клиенты (6 мес)", "👑 VIP Клиенты", "👑 VIP-Клиенты"]))
@router.callback_query(F.data == "admin_vip_clients")
async def msg_admin_vip_clients(event: Message | CallbackQuery, session: AsyncSession):
    """Shows the current members of the closed VIP group (6-month tariff) as a
    list of clickable profile links, plus an Excel export button."""
    now_naive = datetime.datetime.utcnow()
    clients = await _get_active_vip_clients(session)
    answer = event.answer if isinstance(event, Message) else event.message.answer

    if not clients:
        await answer(
            "👑 <b>VIP-Клиенты (6 месяцев)</b>\n\nСейчас нет активных участников 6-месячного тарифа.",
            parse_mode="HTML",
        )
        if isinstance(event, CallbackQuery):
            await event.answer()
        return

    lines = [f"👑 <b>VIP-Клиенты (6 месяцев) — {len(clients)} чел.</b>\n"]
    for idx, (s, u) in enumerate(clients, start=1):
        exp_dt = to_naive_utc(s.expires_at)
        days_left = max(0, (exp_dt - now_naive).days) if exp_dt else 0
        name = escape(u.full_name if u and u.full_name else "Без имени")
        profile_link = f"<a href='tg://user?id={s.user_id}'>{name}</a>"
        uname = f" (@{escape(u.username)})" if (u and u.username) else ""
        line = f"{idx}. {profile_link}{uname} — <code>{s.user_id}</code> — ⏳ {days_left} дн."
        if sum(map(len, lines)) + len(lines) + len(line) > 3800:
            lines.append("\n… (полный список — в Excel)")
            break
        lines.append(line)

    text = "\n".join(lines)

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📥 Выгрузить список в Excel", callback_data="vip_excel_download")]
    ])
    await answer(text, reply_markup=kb, parse_mode="HTML")
    if isinstance(event, CallbackQuery):
        await event.answer()


async def generate_vip_excel(session: AsyncSession) -> BytesIO:
    async with _excel_lock:
        clients = await _get_active_vip_clients(session)
        await session.rollback()
        return await _render_in_thread(_render_vip_excel, clients, datetime.datetime.utcnow())


def _render_vip_excel(
    clients: list[tuple[ExportSubscription, ExportUser]], now: datetime.datetime,
) -> BytesIO:

    header_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="833C0C")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(border_style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "VIP Клиенты (6 мес)"

    headers = ["Telegram ID", "Username", "ФИО", "Телефон", "Дата окончания", "Осталось дней"]
    ws.append(headers)
    ws.row_dimensions[1].height = 25
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = border

    for row_idx, (s, u) in enumerate(clients, start=2):
        exp_dt = to_naive_utc(s.expires_at)
        days_left = max(0, (exp_dt - now).days) if exp_dt else 0
        row = [
            s.user_id,
            f"@{u.username}" if (u and u.username) else "-",
            _excel_text(u.full_name if u else None),
            _excel_text(u.phone_number if u else None),
            exp_dt.strftime("%Y-%m-%d %H:%M") if exp_dt else "-",
            days_left,
        ]
        ws.append(row)
        ws.row_dimensions[row_idx].height = 20
        for col_idx in range(1, len(row) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border
            cell.font = Font(name="Calibri", size=11)
            if col_idx in (1, 5, 6):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(80, max(max_len + 4, 15))

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


@router.callback_query(F.data == "vip_excel_download")
async def cb_vip_excel_download(callback: CallbackQuery, session: AsyncSession, bot: Bot):
    await callback.answer("Генерация Excel…")
    try:
        stream = await generate_vip_excel(session)
        filename = f"vip_clients_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        input_file = BufferedInputFile(stream.getvalue(), filename=filename)
        await bot.send_document(
            chat_id=callback.message.chat.id,
            document=input_file,
            caption="👑 <b>Список VIP-клиентов (6 мес)</b>",
        )
    except Exception as e:
        logger.exception("Error generating VIP Excel")
        await bot.send_message(
            chat_id=callback.message.chat.id,
            text="⚠️ Не удалось сформировать VIP-Excel. Подробности записаны в журнал сервера.",
            parse_mode="HTML",
        )

# ======================== 📊 ПОЛНЫЙ EXCEL-ОТЧЁТ ========================

async def generate_styled_excel(session: AsyncSession) -> BytesIO:
    async with _excel_lock:
        users = [_export_user(user) for user in (await session.scalars(
            select(User).order_by(User.telegram_id))).all()]
        subs = [_export_subscription(sub) for sub in (await session.scalars(
            select(Subscription).order_by(Subscription.expires_at.desc()))).all()]
        payments = [_export_payment(payment) for payment in (await session.scalars(
            select(Payment).order_by(Payment.created_at.desc()))).all()]
        await session.rollback()
        return await _render_in_thread(
            _render_styled_excel, users, subs, payments, datetime.datetime.utcnow(),
        )


def _render_styled_excel(
    users: list[ExportUser], subs: list[ExportSubscription],
    payments: list[ExportPayment], now: datetime.datetime,
) -> BytesIO:
    start_today = datetime.datetime(now.year, now.month, now.day)
    start_month = datetime.datetime(now.year, now.month, 1)

    user_active_sub = {}
    user_latest_sub = {}
    active_vip_subs = []
    
    for s in subs:
        exp_dt = to_naive_utc(s.expires_at)
        if s.user_id not in user_latest_sub:
            user_latest_sub[s.user_id] = s
        if s.status == "active" and exp_dt and exp_dt > now:
            if s.user_id not in user_active_sub:
                user_active_sub[s.user_id] = s
            if s.tariff_months == 6:
                active_vip_subs.append(s)

    rev_today = sum(p.amount for p in payments if p.status == "completed" and to_naive_utc(p.created_at) and to_naive_utc(p.created_at) >= start_today)
    rev_month = sum(p.amount for p in payments if p.status == "completed" and to_naive_utc(p.created_at) and to_naive_utc(p.created_at) >= start_month)
    rev_ltv = sum(p.amount for p in payments if p.status == "completed")
    
    user_ltv = {}
    user_buy_count = {}
    tariff_counts = {1: 0, 3: 0, 6: 0}
    tariff_revs = {1: 0, 3: 0, 6: 0}
    
    for p in payments:
        if p.status == "completed":
            user_ltv[p.user_id] = user_ltv.get(p.user_id, 0) + p.amount
            user_buy_count[p.user_id] = user_buy_count.get(p.user_id, 0) + 1
            t_m = p.tariff_months or 1
            if t_m in tariff_counts:
                tariff_counts[t_m] += 1
                tariff_revs[t_m] += p.amount

    total_users = len(users)
    paying_users = len(user_buy_count)
    repeat_buyers = sum(1 for c in user_buy_count.values() if c > 1)
    retention_rate = (repeat_buyers / paying_users * 100) if paying_users > 0 else 0.0

    user_by_id = {u.telegram_id: u for u in users}

    wb = Workbook()

    # Shared Styles
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    gold_header_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    gold_header_font = Font(name="Calibri", size=11, bold=True, color="833C0C")

    active_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    active_font = Font(name="Calibri", size=11, color="375623")

    expired_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    expired_font = Font(name="Calibri", size=11, color="C65911")

    none_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    none_font = Font(name="Calibri", size=11, color="000000")

    thin_border_side = Side(border_style="thin", color="D9D9D9")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # -------------------------------------------------------------
    # SHEET 1: Сводка и Аналитика
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Сводка и Аналитика"
    ws1.views.sheetView[0].showGridLines = True

    ws1.append(["📊 СВОДНЫЙ ОТЧЁТ И КОГОРТНАЯ АНАЛИТИКА"])
    ws1.cell(row=1, column=1).font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    ws1.append([])

    ws1.append(["Показатель", "Значение"])
    for col_idx in range(1, 3):
        c = ws1.cell(row=3, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = header_align
        c.border = thin_border

    kpis = [
        ("Касса за сегодня (UZS)", rev_today, '#,##0 "UZS"'),
        ("Выручка за текущий месяц (UZS)", rev_month, '#,##0 "UZS"'),
        ("Общая выручка LTV (UZS)", rev_ltv, '#,##0 "UZS"'),
        ("Всего пользователей", total_users, '0'),
        ("Платящих клиентов", paying_users, '0'),
        ("Повторных покупателей", repeat_buyers, '0'),
        ("Коэффициент удержания (Retention)", retention_rate / 100.0, '0.0%')
    ]

    for idx, (label, val, fmt) in enumerate(kpis, start=4):
        ws1.append([label, val])
        c1 = ws1.cell(row=idx, column=1)
        c2 = ws1.cell(row=idx, column=2)
        c1.font = Font(name="Calibri", size=11, bold=True)
        c1.border = thin_border
        c2.font = Font(name="Calibri", size=11)
        c2.border = thin_border
        c2.number_format = fmt
        c2.alignment = Alignment(horizontal="right")

    ws1.append([])
    start_r2 = len(kpis) + 6
    ws1.cell(row=start_r2, column=1, value="Тариф").fill = header_fill
    ws1.cell(row=start_r2, column=1).font = header_font
    ws1.cell(row=start_r2, column=1).border = thin_border
    
    ws1.cell(row=start_r2, column=2, value="Количество продаж").fill = header_fill
    ws1.cell(row=start_r2, column=2).font = header_font
    ws1.cell(row=start_r2, column=2).border = thin_border

    ws1.cell(row=start_r2, column=3, value="Сумма выручки (UZS)").fill = header_fill
    ws1.cell(row=start_r2, column=3).font = header_font
    ws1.cell(row=start_r2, column=3).border = thin_border

    tariffs_info = [
        ("1 месяц", tariff_counts[1], tariff_revs[1]),
        ("3 месяца", tariff_counts[3], tariff_revs[3]),
        ("6 месяцев (VIP)", tariff_counts[6], tariff_revs[6])
    ]

    for offset, (t_lbl, cnt, amt) in enumerate(tariffs_info, start=start_r2+1):
        ws1.append([t_lbl, cnt, amt])
        ws1.cell(row=offset, column=1).font = Font(name="Calibri", size=11)
        ws1.cell(row=offset, column=1).border = thin_border
        
        c2 = ws1.cell(row=offset, column=2)
        c2.font = Font(name="Calibri", size=11)
        c2.border = thin_border
        c2.alignment = Alignment(horizontal="center")
        
        c3 = ws1.cell(row=offset, column=3)
        c3.font = Font(name="Calibri", size=11)
        c3.border = thin_border
        c3.number_format = '#,##0 "UZS"'
        c3.alignment = Alignment(horizontal="right")

    # Table 3: Comparative Monthly Cohort Dynamics
    ws1.append([])
    ws1.append([])
    start_r3 = start_r2 + len(tariffs_info) + 3
    
    ws1.cell(row=start_r3 - 1, column=1, value="📅 СРАВНИТЕЛЬНАЯ ДИНАМИКА ПО МЕСЯЦАМ (КОГОРТЫ)").font = Font(name="Calibri", size=12, bold=True, color="1F4E79")

    dyn_headers = [
        "Месяц (YYYY-MM)", 
        "Новых пользователей", 
        "Тариф 1 мес (шт)", 
        "Тариф 3 мес (шт)", 
        "Тариф 6 мес VIP (шт)", 
        "Количество продлений (шт)", 
        "Выручка за данный месяц (UZS)"
    ]
    
    ws1.append(dyn_headers)
    for col_idx in range(1, len(dyn_headers) + 1):
        cell = ws1.cell(row=start_r3, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border

    months_set = set()
    user_regs_by_month = {}
    for u in users:
        u_dt = to_naive_utc(u.created_at)
        if u_dt:
            m_str = u_dt.strftime("%Y-%m")
            months_set.add(m_str)
            user_regs_by_month[m_str] = user_regs_by_month.get(m_str, 0) + 1

    sorted_completed_payments = sorted(
        [p for p in payments if p.status == "completed"], 
        key=lambda p: to_naive_utc(p.created_at) or datetime.datetime.min
    )

    user_payment_history = {}
    pay_tariffs_by_month = {}
    pay_renewals_by_month = {}
    pay_rev_by_month = {}

    for p in sorted_completed_payments:
        p_dt = to_naive_utc(p.created_at)
        if p_dt:
            m_str = p_dt.strftime("%Y-%m")
            months_set.add(m_str)
            
            if m_str not in pay_tariffs_by_month:
                pay_tariffs_by_month[m_str] = {1: 0, 3: 0, 6: 0}
            t_m = p.tariff_months or 1
            if t_m in pay_tariffs_by_month[m_str]:
                pay_tariffs_by_month[m_str][t_m] += 1
                
            pay_rev_by_month[m_str] = pay_rev_by_month.get(m_str, 0) + p.amount
            
            u_prev = user_payment_history.get(p.user_id, 0)
            if u_prev > 0:
                pay_renewals_by_month[m_str] = pay_renewals_by_month.get(m_str, 0) + 1
            user_payment_history[p.user_id] = u_prev + 1

    all_months = sorted(list(months_set), reverse=True)
    if not all_months:
        all_months = [now.strftime("%Y-%m")]

    for row_offset, m_str in enumerate(all_months, start=start_r3+1):
        regs = user_regs_by_month.get(m_str, 0)
        t_dict = pay_tariffs_by_month.get(m_str, {1: 0, 3: 0, 6: 0})
        t1 = t_dict.get(1, 0)
        t3 = t_dict.get(3, 0)
        t6 = t_dict.get(6, 0)
        ren = pay_renewals_by_month.get(m_str, 0)
        rev = pay_rev_by_month.get(m_str, 0)

        row_vals = [m_str, regs, t1, t3, t6, ren, rev]
        ws1.append(row_vals)
        ws1.row_dimensions[row_offset].height = 20

        for col_idx in range(1, len(row_vals) + 1):
            cell = ws1.cell(row=row_offset, column=col_idx)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=11)
            if col_idx == 7:
                cell.number_format = '#,##0 "UZS"'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = min(80, max(max_len + 4, 18))

    # -------------------------------------------------------------
    # SHEET 2: База пользователей
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="База пользователей")
    ws2.views.sheetView[0].showGridLines = True

    u_headers = [
        "Telegram ID", "Username", "ФИО", "Телефон", "Язык", 
        "Баланс (UZS)", "Приглашённых", "Статус подписки", "Тариф (мес)", 
        "Дата окончания", "Дата регистрации"
    ]
    ws2.append(u_headers)
    ws2.row_dimensions[1].height = 25
    for col_idx in range(1, len(u_headers) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border

    ref_counts = {}
    for u in users:
        if u.referred_by:
            ref_counts[u.referred_by] = ref_counts.get(u.referred_by, 0) + 1

    for row_idx, u in enumerate(users, start=2):
        active_sub = user_active_sub.get(u.telegram_id)
        latest_sub = user_latest_sub.get(u.telegram_id)
        buys = user_buy_count.get(u.telegram_id, 0)

        if active_sub:
            status = "🟢 Активен"
            row_fill = active_fill
            row_font = active_font
            exp_dt = to_naive_utc(active_sub.expires_at)
            exp_str = exp_dt.strftime("%Y-%m-%d %H:%M") if exp_dt else "-"
            t_str = f"{active_sub.tariff_months} мес."
        elif buys > 0 or latest_sub:
            status = "🔴 Истекла"
            row_fill = expired_fill
            row_font = expired_font
            exp_dt = to_naive_utc(latest_sub.expires_at) if latest_sub else None
            exp_str = exp_dt.strftime("%Y-%m-%d %H:%M") if exp_dt else "-"
            t_str = f"{latest_sub.tariff_months} мес." if latest_sub else "-"
        else:
            status = "⚪️ Без покупок"
            row_fill = none_fill
            row_font = none_font
            exp_str = "-"
            t_str = "-"

        u_reg_dt = to_naive_utc(u.created_at)
        reg_str = u_reg_dt.strftime("%Y-%m-%d %H:%M") if u_reg_dt else "-"

        row_data = [
            u.telegram_id,
            f"@{u.username}" if u.username else "-",
            _excel_text(u.full_name),
            _excel_text(u.phone_number),
            (u.language or "uz").upper(),
            u.balance or 0,
            ref_counts.get(u.telegram_id, 0),
            status,
            t_str,
            exp_str,
            reg_str
        ]
        ws2.append(row_data)
        ws2.row_dimensions[row_idx].height = 20

        for col_idx in range(1, len(row_data) + 1):
            cell = ws2.cell(row=row_idx, column=col_idx)
            cell.fill = row_fill
            cell.font = row_font
            cell.border = thin_border
            if col_idx == 6:
                cell.number_format = '#,##0 "UZS"'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in (1, 5, 7, 8, 9, 10, 11):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = min(80, max(max_len + 4, 14))

    # -------------------------------------------------------------
    # SHEET 3: VIP Клиенты (6 мес)
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="VIP Клиенты (6 мес)")
    ws3.views.sheetView[0].showGridLines = True

    vip_headers = [
        "Telegram ID", "Username", "ФИО", "Телефон", "Дата окончания",
        "Осталось дней"
    ]
    ws3.append(vip_headers)
    ws3.row_dimensions[1].height = 25
    for col_idx in range(1, len(vip_headers) + 1):
        cell = ws3.cell(row=1, column=col_idx)
        cell.fill = gold_header_fill
        cell.font = gold_header_font
        cell.alignment = header_align
        cell.border = thin_border

    for row_idx, s in enumerate(active_vip_subs, start=2):
        u = user_by_id.get(s.user_id)
        exp_dt = to_naive_utc(s.expires_at)
        days_left = max(0, (exp_dt - now).days) if exp_dt else 0

        row_data = [
            s.user_id,
            f"@{u.username}" if u and u.username else "-",
            _excel_text(u.full_name if u else None),
            _excel_text(u.phone_number if u else None),
            exp_dt.strftime("%Y-%m-%d %H:%M") if exp_dt else "-",
            days_left
        ]
        ws3.append(row_data)
        ws3.row_dimensions[row_idx].height = 20

        for col_idx in range(1, len(row_data) + 1):
            cell = ws3.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=11)
            if col_idx in (1, 5, 6):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws3.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = min(80, max(max_len + 4, 15))

    # -------------------------------------------------------------
    # SHEET 4: Транзакции
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="Транзакции")
    ws4.views.sheetView[0].showGridLines = True

    tx_headers = [
        "ID Платёжа", "Telegram ID", "ФИО", "Тариф (мес)", 
        "Сумма тарифа (UZS)", "Списано кешбэка (UZS)", "Итого оплачено (UZS)", 
        "Статус", "Дата платежа"
    ]
    ws4.append(tx_headers)
    ws4.row_dimensions[1].height = 25
    for col_idx in range(1, len(tx_headers) + 1):
        cell = ws4.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border

    for row_idx, p in enumerate(payments, start=2):
        u = user_by_id.get(p.user_id)
        st_icon = "✅ Успешно" if p.status == "completed" else "⏳ В обработке" if p.status == "pending" else "❌ Ошибка"
        p_dt = to_naive_utc(p.created_at)

        cashback_applied = p.cashback_applied or 0
        paid_amount = p.amount or 0
        tariff_price = paid_amount + cashback_applied
        row_data = [
            p.id,
            p.user_id,
            _excel_text(u.full_name if u else None),
            f"{p.tariff_months} мес.",
            tariff_price,
            cashback_applied,
            paid_amount,
            st_icon,
            p_dt.strftime("%Y-%m-%d %H:%M") if p_dt else "-"
        ]
        ws4.append(row_data)
        ws4.row_dimensions[row_idx].height = 20

        for col_idx in range(1, len(row_data) + 1):
            cell = ws4.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=11)
            if col_idx in (5, 6, 7):
                cell.number_format = '#,##0 "UZS"'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in (1, 2, 4, 8, 9):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws4.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws4.column_dimensions[col_letter].width = min(80, max(max_len + 4, 15))

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

@router.message(F.text.in_([
    "👥 Пользователи (База Excel)",
    "👥 Пользователи (Выгрузить Excel)", 
    "👥 Пользователи", 
    "📥 Скачать ПОЛНЫЙ Excel",
    "📥 Скачать полный Excel", 
    "📊 Скачать полный отчёт (Excel)", 
    "📥 Скачать Excel", 
    "📊 Скачать Excel"
]))
@router.callback_query(F.data == "users_excel_download")
async def msg_an_report_dl(event: Message | CallbackQuery, session: AsyncSession):
    bot = event.bot if isinstance(event, Message) else event.message.bot
    chat_id = event.chat.id if isinstance(event, Message) else event.message.chat.id

    if isinstance(event, CallbackQuery):
        await event.answer("Генерация Excel...", show_alert=False)
        wait_msg = await event.message.answer("📥 <i>Генерация полного Excel-отчета...</i>")
    else:
        wait_msg = await event.answer("📥 <i>Генерация полного Excel-отчета...</i>")

    try:
        excel_stream = await generate_styled_excel(session)
        filename = f"full_admin_report_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        input_file = BufferedInputFile(excel_stream.getvalue(), filename=filename)
        
        try:
            await wait_msg.delete()
        except Exception:
            pass
            
        await bot.send_document(
            chat_id=chat_id, 
            document=input_file, 
            caption="📊 <b>Полный сводный Excel-отчёт</b>\n(Сводка и Аналитика, База пользователей, VIP Клиенты (6 мес), Транзакции)"
        )
    except Exception as e:
        logger.exception("Error generating or sending Excel report")
        try:
            await wait_msg.delete()
        except Exception:
            pass
        err_msg = (
            "⚠️ <b>Ошибка при генерации Excel.</b>\n\n"
            f"Пожалуйста, проверьте логи сервера."
        )
        await bot.send_message(chat_id=chat_id, text=err_msg, parse_mode="HTML")

# ======================== 📢 РАССЫЛКА СООБЩЕНИЙ ========================

@router.message(F.text == "📢 Рассылка сообщений")
async def msg_admin_broadcast(message: Message, state: FSMContext):
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="👥 Всем пользователям", callback_data="bc_aud_all")],
        [InlineKeyboardButton(text="🟢 Активным подписчикам", callback_data="bc_aud_active")],
        [InlineKeyboardButton(text="🔴 Истекшим подпискам", callback_data="bc_aud_expired")],
        [InlineKeyboardButton(text="⚪️ Пользователям без покупок", callback_data="bc_aud_never")],
        [InlineKeyboardButton(text="◀️ Отмена", callback_data="admin_main")]
    ])
    await message.answer("📢 <b>Выберите аудиторию для рассылки:</b>", reply_markup=keyboard)
    await state.set_state(AdminBroadcast.waiting_for_audience)

@router.callback_query(AdminBroadcast.waiting_for_audience, F.data.in_(["bc_aud_all", "bc_aud_active", "bc_aud_expired", "bc_aud_never"]))
async def process_bc_audience(callback: CallbackQuery, state: FSMContext):
    aud = callback.data.replace("bc_aud_", "")
    await state.update_data(audience=aud)
    await state.set_state(AdminBroadcast.waiting_for_message)
    await callback.message.edit_text(
        "📝 <b>Отправьте сообщение для рассылки.</b>\n\n"
        "Вы можете отправить текст, фото, видео или аудио. Оно будет переслано пользователям в точности так, как вы его отправите.",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="◀️ Отмена", callback_data="admin_main")]])
    )
    await callback.answer()

@router.message(AdminBroadcast.waiting_for_message)
async def process_bc_message(message: Message, state: FSMContext):
    await state.update_data(msg_id=message.message_id, chat_id=message.chat.id)
    await state.set_state(AdminBroadcast.waiting_for_confirm)
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🚀 Запустить рассылку", callback_data="bc_start_now")],
        [InlineKeyboardButton(text="❌ Отменить", callback_data="admin_main")]
    ])
    await message.answer("⚠️ <b>Подтвердите запуск рассылки.</b>\nСообщение готово к отправке.", reply_markup=keyboard)

@router.callback_query(AdminBroadcast.waiting_for_confirm, F.data == "bc_start_now")
async def process_bc_start(callback: CallbackQuery, state: FSMContext, session: AsyncSession, bot: Bot):
    data = await state.get_data()
    audience = data.get("audience")
    if audience not in {"all", "active", "expired", "never"} or not all(
        isinstance(data.get(key), int) for key in ("chat_id", "msg_id")
    ):
        await state.clear()
        await callback.answer("Данные рассылки устарели. Создайте её заново.", show_alert=True)
        return
    # Drop the confirmation state before I/O; per-user FSM isolation in runtime
    # serializes repeat callbacks for the same administrator.
    await state.clear()
    now = datetime.datetime.now(datetime.timezone.utc)
    active = exists().where(Subscription.user_id == User.telegram_id,
                            Subscription.status == "active", Subscription.expires_at > now)
    ever_subscribed = exists().where(Subscription.user_id == User.telegram_id)
    ever_paid = exists().where(Payment.user_id == User.telegram_id, Payment.status == "completed")
    query = select(User.telegram_id).order_by(User.telegram_id)
    if audience == "active":
        query = query.where(active)
    elif audience == "expired":
        query = query.where(~active, ever_subscribed | ever_paid)
    elif audience == "never":
        query = query.where(~ever_subscribed, ~ever_paid)
    user_ids = list((await session.scalars(query)).all())
    await session.rollback()  # release the read transaction before long network work
    await callback.message.edit_text(f"🚀 <b>Рассылка запущена!</b> Ожидайте окончания...\nАудитория: {len(user_ids)} чел.")
    
    async def do_broadcast(uids: list[int], from_chat: int, msg_id: int) -> None:
        sent = 0
        for uid in uids:
            try:
                await bot.copy_message(chat_id=uid, from_chat_id=from_chat, message_id=msg_id)
                sent += 1
            except Exception:
                logger.warning("Broadcast delivery failed for user %s", uid, exc_info=True)
            finally:
                # Also pace failures; the runtime limiter handles global 429s.
                await asyncio.sleep(0.05)
        try:
            await bot.send_message(chat_id=callback.from_user.id, text=f"✅ <b>Рассылка завершена!</b>\nУспешно доставлено: {sent} / {len(uids)}.")
        except Exception:
            pass

    create_background_task(do_broadcast(user_ids, data['chat_id'], data['msg_id']), name="admin-broadcast")
    await state.clear()
    await callback.answer()

# ======================== 📥 ЗАЯВКИ И ТИКЕТЫ ========================

@router.message(F.text.startswith("📥 Заявки"))
async def msg_admin_tickets(message: Message, session: AsyncSession):
    stmt = select(Ticket, User).join(User, Ticket.user_id == User.telegram_id).where(Ticket.status == "open").limit(10)
    result = await session.execute(stmt)
    tickets = result.all()
    
    if not tickets:
        await message.answer("✅ Открытых заявок/тикетов нет.")
        return
        
    for ticket, user in tickets:
        text = f"🎫 Тикет #{ticket.id}\n👤 Пользователь: {escape(user.full_name or 'Без имени')} (<code>{user.telegram_id}</code>)\n\n📝 <i>{escape((ticket.text or '')[:2500])}</i>"
        keyboard = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="💬 Ответить", callback_data=f"adm_ticket_reply_{ticket.id}")]
        ])
        await message.answer(text, reply_markup=keyboard)

@router.callback_query(F.data.startswith("adm_ticket_reply_"))
async def adm_ticket_reply_prompt(callback: CallbackQuery, state: FSMContext):
    ticket_id = _positive_id((callback.data or "").removeprefix("adm_ticket_reply_"))
    if ticket_id is None:
        await callback.answer("Некорректный номер тикета.", show_alert=True)
        return
    await state.set_state(AdminTicketState.waiting_for_reply)
    await state.update_data(ticket_id=ticket_id)
    await callback.message.answer(f"📝 Отправьте ваш ответ на тикет #{ticket_id}:")
    await callback.answer()

@router.message(AdminTicketState.waiting_for_reply, F.text)
async def process_adm_ticket_reply(message: Message, state: FSMContext, session: AsyncSession, bot: Bot):
    data = await state.get_data()
    ticket_id = data.get("ticket_id")
    
    ticket = await session.scalar(select(Ticket).where(Ticket.id == ticket_id))
    if ticket:
        ticket.status = "closed"
        await session.commit()
        
        user = await session.scalar(select(User).where(User.telegram_id == ticket.user_id))
        lang = user.language if user and user.language else "uz"
        
        reply_msg = f"📩 <b>Ответ поддержки по тикету #{ticket_id}:</b>\n\n{escape(message.text)}"
        try:
            await bot.send_message(chat_id=ticket.user_id, text=reply_msg)
            await message.answer(f"✅ Ответ отправлен пользователю {ticket.user_id}, тикет #{ticket_id} закрыт.")
        except Exception as e:
            await message.answer(f"⚠️ Тикет закрыт, но сообщение пользователю не доставлено: {e}")
    else:
        await message.answer("❌ Тикет не найден.")

    await state.clear()

# ======================== 💳 ПОДТВЕРЖДЕНИЕ / ОТКЛОНЕНИЕ ОПЛАТ ========================

@router.callback_query(F.data.startswith("admin_conf_"))
async def admin_confirm_payment(callback: CallbackQuery, session: AsyncSession, bot: Bot):
    """Confirms a pending payment: activates/extends subscription, deducts cashback,
    awards referral bonus, issues single-use invite link and notifies the user.
    Delegates to the shared payment service so all money logic lives in one place."""
    try:
        _, _, payment_id_str, _user_id_str = (callback.data or "").split("_")
        if _positive_id(_user_id_str) is None:
            raise ValueError("invalid user ID")
        payment_id = _positive_id(payment_id_str)
        if payment_id is None:
            raise ValueError("invalid payment ID")
    except (ValueError, IndexError):
        await callback.answer("❌ Некорректные данные заявки.", show_alert=True)
        return

    try:
        ok = await process_successful_payment(payment_id, session, bot)
    except Exception as e:
        logger.exception(f"Error confirming payment {payment_id}")
        await session.rollback()
        await callback.answer("❌ Не удалось подтвердить платёж. Проверьте заявку и журнал сервера.", show_alert=True)
        return

    if ok:
        try:
            await callback.message.edit_text(
                callback.message.html_text + "\n\n✅ <b>Оплата подтверждена, подписка активирована.</b>"
            )
        except Exception:
            pass
        await callback.answer("✅ Оплата подтверждена!")
    else:
        try:
            await callback.message.edit_text(
                callback.message.html_text + "\n\n⚠️ <b>Заявка не найдена или уже обработана.</b>"
            )
        except Exception:
            pass
        await callback.answer("Заявка не найдена или уже обработана.", show_alert=True)


@router.callback_query(F.data.startswith("admin_rej_"))
async def admin_reject_payment(callback: CallbackQuery, session: AsyncSession, bot: Bot):
    """Rejects a pending payment and notifies the user."""
    try:
        _, _, payment_id_str, user_id_str = (callback.data or "").split("_")
        payment_id = _positive_id(payment_id_str)
        if payment_id is None:
            raise ValueError("invalid payment ID")
        if _positive_id(user_id_str) is None:
            raise ValueError("invalid user ID")
    except (ValueError, IndexError):
        await callback.answer("❌ Некорректные данные заявки.", show_alert=True)
        return

    try:
        user_id = await reject_pending_payment(payment_id, session)
    except Exception:
        await session.rollback()
        logger.exception("Payment rejection failed for %s", payment_id)
        await callback.answer("Не удалось отклонить платёж. Повторите позже.", show_alert=True)
        return
    if user_id is None:
        try:
            await callback.message.edit_text(
                callback.message.html_text + "\n\n⚠️ <b>Заявка не найдена или уже обработана.</b>"
            )
        except Exception:
            pass
        await callback.answer("Заявка не найдена или уже обработана.", show_alert=True)
        return

    user = await session.scalar(select(User).where(User.telegram_id == user_id))
    lang = user.language if user and user.language else "uz"
    try:
        await callback.message.edit_text(callback.message.html_text + "\n\n❌ <b>Оплата отклонена.</b>")
    except Exception:
        pass
    try:
        await bot.send_message(chat_id=user_id, text=texts.PAYMENT_FAILED_NOTIFY[lang])
    except Exception as e:
        logger.error(f"Failed to notify user {user_id} about payment rejection: {e}")
    await callback.answer("❌ Оплата отклонена.")

# ======================== 🔎 ПОИСК И УПРАВЛЕНИЕ КЛИЕНТОМ ========================

@router.message(F.text.in_(["🔎 Поиск клиента", "👥 Поиск клиента"]))
async def msg_search_user(message: Message, state: FSMContext):
    await message.answer(
        "🔍 <b>Найти клиента</b>\n\n"
        "Отправьте Telegram ID клиента или его @username (можно без «@»).",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="◀️ Отмена", callback_data="admin_main")]])
    )
    await state.set_state(AdminUserSearch.waiting_for_query)


@router.message(AdminUserSearch.waiting_for_query, F.text)
async def process_user_search(message: Message, state: FSMContext, session: AsyncSession):
    query = (message.text or "").strip()
    if not query or len(query) > 64:
        await message.answer("Укажите Telegram ID или username длиной до 64 символов.")
        return

    if query.isdecimal():
        user_id = _positive_id(query)
        if user_id is None:
            await message.answer("Некорректный Telegram ID.")
            return
        stmt = select(User).where(User.telegram_id == user_id)
    else:
        query = query.removeprefix("@")
        if not query or not query.isascii() or not all(c.isalnum() or c == "_" for c in query):
            await message.answer("Username должен содержать латинские буквы, цифры или _. ")
            return
        stmt = select(User).where(func.lower(User.username) == query.lower())

    user = await session.scalar(stmt)

    if not user:
        await message.answer("❌ Пользователь не найден. Попробуйте ещё раз или нажмите /admin")
        return

    sub = await session.scalar(
        select(Subscription).where(
            Subscription.user_id == user.telegram_id,
            Subscription.status == "active"
        ).order_by(Subscription.expires_at.desc()).limit(1)
    )

    now = datetime.datetime.utcnow()
    if sub and to_naive_utc(sub.expires_at) and to_naive_utc(sub.expires_at) > now:
        status = "✅ АКТИВНА"
        expires = sub.expires_at.strftime("%Y-%m-%d %H:%M UTC")
        tariff = f"{sub.tariff_months} мес."
    else:
        status = "❌ ИСТЕКЛА / НЕТ"
        expires = "-"
        tariff = "-"

    text = (
        f"👤 <b>Карточка клиента</b>\n\n"
        f"ID: <code>{user.telegram_id}</code>\n"
        f"Username: {escape('@' + user.username) if user.username else 'Нет'}\n"
        f"Имя: {escape(user.full_name or 'Нет')}\n"
        f"Телефон: <code>{escape(user.phone_number or 'Нет')}</code>\n"
        f"Баланс: {user.balance or 0:,} UZS\n\n"
        f"<b>Подписка:</b> {status}\n"
        f"<b>Тариф:</b> {tariff}\n"
        f"<b>Истекает:</b> {expires}\n"
    )

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="➕ Продлить (30 дней)", callback_data=f"adm_user_ext_{user.telegram_id}")],
        [InlineKeyboardButton(text="⛔️ Аннулировать (Кик)", callback_data=f"adm_user_kick_{user.telegram_id}")],
        [InlineKeyboardButton(text="🔗 Сгенерировать ссылку", callback_data=f"adm_user_link_{user.telegram_id}")],
        [InlineKeyboardButton(text="👤 Открыть профиль", url=f"tg://user?id={user.telegram_id}")],
        [InlineKeyboardButton(text="◀️ В меню", callback_data="admin_main")]
    ])

    await message.answer(text, reply_markup=keyboard)
    await state.clear()


@router.callback_query(F.data.startswith("adm_user_ext_"))
async def adm_user_extend(callback: CallbackQuery, session: AsyncSession) -> None:
    user_id = _positive_id((callback.data or "").removeprefix("adm_user_ext_"))
    if user_id is None:
        await callback.answer("Некорректный ID клиента.", show_alert=True)
        return
    try:
        user = await session.scalar(select(User).where(User.telegram_id == user_id)
                                    .with_for_update().execution_options(populate_existing=True))
        if user is None:
            await session.rollback()
            await callback.answer("Клиент не найден.", show_alert=True)
            return
        sub = await session.scalar(
            select(Subscription).where(Subscription.user_id == user_id, Subscription.status == "active")
            .order_by(Subscription.expires_at.desc()).limit(1)
            .with_for_update().execution_options(populate_existing=True)
        )
        now = datetime.datetime.now(datetime.timezone.utc)
        expiry = to_naive_utc(sub.expires_at) if sub else None
        started_at = expiry.replace(tzinfo=datetime.timezone.utc) if expiry and expiry > to_naive_utc(now) else now
        expires_at = started_at + datetime.timedelta(days=30)
        if sub:
            sub.expires_at = expires_at
            sub.notified_3d = False
            sub.notified_1d = False
        else:
            session.add(Subscription(user_id=user_id, status="active", tariff_months=1,
                                     started_at=started_at, expires_at=expires_at))
        await session.commit()
    except Exception:
        await session.rollback()
        logger.exception("Failed to extend subscription for user %s", user_id)
        await callback.answer("Не удалось продлить подписку. Повторите позже.", show_alert=True)
        return
    await callback.answer("✅ Подписка продлена на 30 дней!", show_alert=True)
    await callback.message.edit_reply_markup(
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="◀️ В меню", callback_data="admin_main")]])
    )


@router.callback_query(F.data.startswith("adm_user_kick_"))
async def adm_user_kick(callback: CallbackQuery, session: AsyncSession, bot: Bot) -> None:
    user_id = _positive_id((callback.data or "").removeprefix("adm_user_kick_"))
    if user_id is None:
        await callback.answer("Некорректный ID клиента.", show_alert=True)
        return
    try:
        user = await session.scalar(select(User).where(User.telegram_id == user_id)
                                    .with_for_update().execution_options(populate_existing=True))
        if user is None:
            await session.rollback()
            await callback.answer("Клиент не найден.", show_alert=True)
            return
        subs = (await session.scalars(
            select(Subscription).where(Subscription.user_id == user_id,
                                       Subscription.status.in_(["active", "revocation_pending"]))
            .with_for_update().execution_options(populate_existing=True)
        )).all()
        for sub in subs:
            sub.status = "revocation_pending"
        await session.commit()  # durable retry marker before any Telegram calls

        await session.scalar(select(User).where(User.telegram_id == user_id)
                             .with_for_update().execution_options(populate_existing=True))
        # A payment may have activated a new subscription after the marker commit.
        current = (await session.scalars(
            select(Subscription).where(Subscription.user_id == user_id)
            .with_for_update().execution_options(populate_existing=True)
        )).all()
        now = datetime.datetime.utcnow()
        active = [sub for sub in current if sub.status == "active"
                  and to_naive_utc(sub.expires_at) and to_naive_utc(sub.expires_at) > now]
        pending = [sub for sub in current if sub.status == "revocation_pending"]
        for sub in pending:
            if config.channel_id and sub.invite_link:
                await revoke_invite(bot, config.channel_id, sub.invite_link)
            if config.vip_chat_id and sub.vip_invite_link:
                await revoke_invite(bot, config.vip_chat_id, sub.vip_invite_link)
        if pending and not active and config.channel_id:
            await kick_member(bot, config.channel_id, user_id)
        if any(sub.tariff_months == 6 for sub in pending) and not any(sub.tariff_months == 6 for sub in active) and config.vip_chat_id:
            await kick_member(bot, config.vip_chat_id, user_id)
        for sub in pending:
            sub.invite_link = None
            sub.vip_invite_link = None
            sub.status = "expired"
        await session.commit()
    except Exception:
        await session.rollback()
        logger.exception("Admin revocation pending retry for user %s", user_id)
        await callback.answer("Подписка аннулирована. Исключение из чатов будет повторено планировщиком.", show_alert=True)
        return
    await callback.answer("⛔️ Подписка аннулирована; доступ в чаты проверен.", show_alert=True)
    await callback.message.edit_reply_markup(
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="◀️ В меню", callback_data="admin_main")]])
    )


@router.callback_query(F.data.startswith("adm_user_link_"))
async def adm_user_link(callback: CallbackQuery, session: AsyncSession, bot: Bot) -> None:
    user_id = _positive_id((callback.data or "").removeprefix("adm_user_link_"))
    if user_id is None:
        await callback.answer("Некорректный ID клиента.", show_alert=True)
        return
    try:
        user = await session.scalar(select(User).where(User.telegram_id == user_id)
                                    .with_for_update().execution_options(populate_existing=True))
        sub = await session.scalar(
            select(Subscription).where(Subscription.user_id == user_id, Subscription.status == "active",
                                        Subscription.expires_at > datetime.datetime.now(datetime.timezone.utc))
            .order_by(Subscription.expires_at.desc()).limit(1)
            .with_for_update().execution_options(populate_existing=True)
        ) if user else None
        if sub is None or not config.channel_id:
            await session.rollback()
            await callback.answer("Нет действующей подписки или канал не настроен.", show_alert=True)
            return
        if sub.invite_link:
            await revoke_invite(bot, config.channel_id, sub.invite_link)
        invite = await bot.create_chat_invite_link(
            chat_id=config.channel_id, creates_join_request=True,
            expire_date=sub.expires_at, name=f"Manual: {user_id}", request_timeout=15,
        )
        sub.invite_link = invite.invite_link
        await session.commit()
    except Exception:
        await session.rollback()
        logger.exception("Failed to issue admin invite for user %s", user_id)
        await callback.answer("Не удалось создать ссылку. Повторите позже.", show_alert=True)
        return
    await callback.message.answer(
        f"🔗 <b>Ссылка для ID {user_id}:</b>\n{escape(invite.invite_link)}"
    )
    await callback.answer()
